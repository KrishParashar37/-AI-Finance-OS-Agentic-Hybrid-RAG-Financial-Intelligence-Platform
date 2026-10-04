"""Report builder — PDF, CSV, Excel generation."""
from __future__ import annotations
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from fpdf import FPDF


def _money(v: float) -> str:
    return f"Rs.{v:,.0f}"


def _fmt(d: datetime) -> str:
    return d.strftime("%d %b %Y")


def build_report(report_type: str, from_dt: datetime, to_dt: datetime, txns: list, budgets: list) -> dict:
    end = to_dt.replace(hour=23, minute=59, second=59)
    period_txns = [t for t in txns if from_dt <= t.date <= end]
    exp = [t for t in period_txns if t.type == "expense"]
    inc = [t for t in period_txns if t.type == "income"]
    t_exp = sum(t.amount for t in exp)
    t_inc = sum(t.amount for t in inc)
    period = f"{_fmt(from_dt)} - {_fmt(end)}"
    tx_headers = ["Date", "Merchant", "Category", "Type", "Payment", "Amount", "Tax"]

    def tx_row(t):
        return [_fmt(t.date), t.merchant, t.category, t.type, t.payment_method, str(round(t.amount)), str(round(t.tax))]

    base_summary = [f"Period: {period}", f"Total income: {_money(t_inc)}", f"Total expenses: {_money(t_exp)}", f"Net savings: {_money(t_inc - t_exp)}"]

    if report_type == "expenses":
        return {"title": "Expense Report", "headers": tx_headers, "rows": [tx_row(t) for t in exp], "summary": [f"Period: {period}", f"Total: {_money(t_exp)}", f"Transactions: {len(exp)}"]}
    if report_type == "income":
        return {"title": "Income Report", "headers": tx_headers, "rows": [tx_row(t) for t in inc], "summary": [f"Period: {period}", f"Total: {_money(t_inc)}", f"Transactions: {len(inc)}"]}
    if report_type == "tax":
        taxed = [t for t in exp if t.tax > 0]
        return {"title": "Tax Report (GST paid)", "headers": tx_headers, "rows": [tx_row(t) for t in taxed], "summary": [f"Period: {period}", f"Total GST: {_money(sum(t.tax for t in taxed))}", f"Taxable purchases: {len(taxed)}"]}
    if report_type == "budget":
        from analytics import category_stats
        cats = {c["name"]: c["mtd"] for c in category_stats(txns)}
        from_d, to_d = from_dt, to_dt
        months = max(1, round((to_d - from_d).days / 30))
        rows = [[b.category, str(round(b.limit_amount * months)), str(round(cats.get(b.category, 0))), str(round(b.limit_amount * months - cats.get(b.category, 0))), f"{round(cats.get(b.category,0)/b.limit_amount*100) if b.limit_amount else 0}%"] for b in budgets]
        return {"title": "Budget Report", "headers": ["Category", "Budget", "Spent", "Remaining", "Used"], "rows": rows, "summary": base_summary}
    if report_type in ("monthly", "annual"):
        from collections import defaultdict
        buckets: dict[str, dict] = defaultdict(lambda: {"inc": 0.0, "exp": 0.0})
        for t in period_txns:
            k = t.date.strftime("%b %Y") if report_type == "annual" else _fmt(t.date)
            if t.type == "income":
                buckets[k]["inc"] += t.amount
            else:
                buckets[k]["exp"] += t.amount
        rows = [[k, str(round(v["inc"])), str(round(v["exp"])), str(round(v["inc"] - v["exp"]))] for k, v in buckets.items()]
        return {"title": "Monthly Financial Report" if report_type == "monthly" else "Annual Financial Report", "headers": ["Date", "Income", "Expenses", "Net"], "rows": rows, "summary": base_summary}

    # AI Financial / default
    from collections import defaultdict
    cat_totals: dict[str, float] = defaultdict(float)
    for t in exp:
        cat_totals[t.category] += t.amount
    rows = sorted([[c, str(round(v)), f"{round(v/t_exp*100) if t_exp else 0}%"] for c, v in cat_totals.items()], key=lambda r: -float(r[1]))
    return {"title": "AI Financial Report", "headers": ["Category", "Spent", "Share"], "rows": rows, "summary": base_summary}


def to_csv(data: dict) -> str:
    def esc(v: str) -> str:
        return f'"{v.replace(chr(34), chr(34)*2)}"' if any(c in v for c in '",\n') else v
    lines = [",".join(esc(h) for h in data["headers"])]
    for row in data["rows"]:
        lines.append(",".join(esc(str(c)) for c in row))
    return "\n".join(lines)


def to_excel(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = data.get("title", "Report")[:31]
    # Summary sheet
    ws_sum = wb.create_sheet("Summary")
    for line in data.get("summary", []):
        ws_sum.append([line])
    # Data sheet
    ws.append(data["headers"])
    for row in data["rows"]:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(data: dict) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    title = data.get("title", "Report").encode("latin-1", "replace").decode("latin-1")
    pdf.cell(0, 10, title, ln=True)
    pdf.set_font("Helvetica", "", 9)
    for line in data.get("summary", []):
        clean = line.encode("latin-1", "replace").decode("latin-1")
        pdf.cell(0, 6, clean, ln=True)
    pdf.ln(4)
    # Table header
    pdf.set_font("Helvetica", "B", 8)
    col_w = min(40, 190 // max(len(data["headers"]), 1))
    for h in data["headers"]:
        pdf.cell(col_w, 7, h[:18], border=1)
    pdf.ln()
    pdf.set_font("Helvetica", "", 8)
    for row in data["rows"][:500]:
        for cell in row:
            clean = str(cell).encode("latin-1", "replace").decode("latin-1")[:18]
            pdf.cell(col_w, 6, clean, border=1)
        pdf.ln()
    return bytes(pdf.output())
